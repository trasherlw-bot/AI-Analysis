#!/usr/bin/env python3
"""
Выгрузка заявок HelpDeskEddy по номерам из Excel, обезличивание, разбор на
запросы клиента / ответы ИИ («Курьер Сервис Экспресс») и СХЛОПЫВАНИЕ заявок
с одним id переписки (КСЭ..., b2b..., drr...) в одну строку.

Установка:  pip install requests openpyxl
Опционально (лучше вычищает ФИО):  pip install natasha

Запуск:
    python hde_ai_analysis.py                 # обработать файл из папки
    python hde_ai_analysis.py --debug 12345   # показать сырой JSON одной заявки
"""
import collections
import html
import json
import os
import re
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

import hde_rules

# ===================== НАСТРОЙКИ =====================
HDE_DOMAIN = os.getenv("HDE_DOMAIN", "kse")      # yourcompany.helpdeskeddy.com
HDE_EMAIL = os.getenv("HDE_EMAIL", "")          # email пользователя с доступом к API
HDE_API_KEY = os.getenv("HDE_API_KEY", "")      # ключ API
INPUT_FOLDER = Path(os.getenv("DATA_DIR", "."))   # папка, куда кладёте Excel
ID_COLUMN = "Заявка"                            # заголовок колонки с номерами заявок
AI_MARKER = "Курьер Сервис Экспресс"            # подпись ИИ (запасной признак)
AI_USER_IDS = {"301206"}                        # user_id ИИ-агента в HDE (главный признак)
ONLY_AI_REPLIES = True                          # False = брать все ответы поддержки
MASK_LONG_NUMBERS = True                        # маскировать номера накладных/заказов (8+ цифр)
KEY_PREFIXES = ("КСЭ", "b2b", "drr")            # с чего начинается id переписки
PAUSE_SEC = 0.3                                 # пауза между запросами
SAVE_EVERY = 20                                 # сохранять кэш каждые N заявок
OUT_COLS = ["ID переписки", "Запрос клиента", "Ответы ИИ", "Диалог (размечено)"]
# Колонки прошлых версий: при обработке файла удаляются, если они там есть
REMOVED_COLS = ["Заявки в группе", "Кол-во ответов ИИ", "Статус обработки"]
MSK = timezone(timedelta(hours=3))              # время в HDE и в именах файлов — московское

# Лимит API HelpDeskEddy: 300 запросов в минуту на аккаунт. При превышении аккаунт блокируется на 20 минут.
RATE_LIMIT_RPM = int(os.getenv("HDE_RATE_LIMIT_RPM", "300"))
RATE_SAFETY = float(os.getenv("HDE_RATE_SAFETY", "0.9"))   # скрипт держится на 90% лимита
BLOCK_MINUTES = 20
# =====================================================

BASE = os.getenv("HDE_BASE_URL") or f"https://{HDE_DOMAIN}.helpdeskeddy.com/api/v2"
DATE_RE = re.compile(r"(?=\[\d{1,2}:\d{2} \d{2}\.\d{2}\.\d{4}\])")
DATE_PARSE = re.compile(r"\[(\d{1,2}):(\d{2}) (\d{2})\.(\d{2})\.(\d{4})\]")
CELL_LIMIT = 32000  # лимит Excel 32767 символов в ячейке

_pref = "|".join(re.escape(p) for p in KEY_PREFIXES)
# "id":"КСЭ4042192" (в т.ч. если JSON вложен в текст и кавычки экранированы)
KEY_RE_ID = re.compile(r'\\*"id\\*"\s*:\s*\\*"((?:%s)[^"\\\s]*)' % _pref, re.I)
# запасной вариант: просто токен с нужным префиксом и цифрой
KEY_RE_FALLBACK = re.compile(r'(?<![\w])((?:%s)[\w-]*\d[\w-]*)' % _pref, re.I)

# ---------- обезличивание ----------
RE_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
RE_URL = re.compile(r"https?://\S+")
RE_PHONE = re.compile(r"(?<!\d)(?:\+7|8)[\s\-()]*\d{3}[\s\-()]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}(?!\d)")
RE_LONG = re.compile(r"(?<!\d)\d{8,}(?!\d)")

_nlp = None
try:
    from natasha import Doc, NewsEmbedding, NewsNERTagger, Segmenter
    _nlp = (Segmenter(), NewsNERTagger(NewsEmbedding()))
except Exception:
    pass  # без natasha ФИО не маскируются — проверьте выборочно глазами


def anonymize(text: str) -> str:
    text = RE_EMAIL.sub("[EMAIL]", text)
    text = RE_URL.sub("[URL]", text)
    text = RE_PHONE.sub("[ТЕЛЕФОН]", text)
    if MASK_LONG_NUMBERS:
        text = RE_LONG.sub("[НОМЕР]", text)
    if _nlp:
        segmenter, ner = _nlp
        doc = Doc(text)
        doc.segment(segmenter)
        doc.tag_ner(ner)
        for span in sorted(doc.spans, key=lambda s: s.start, reverse=True):
            if span.type == "PER":
                text = text[:span.start] + "[ФИО]" + text[span.stop:]
    return text


def clean_html(s: str) -> str:
    s = re.sub(r"<br\s*/?>", "\n", s, flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)
    s = re.sub(r"[ \t]*\n\s*", "\n", html.unescape(s))   # убрать лишние пустые строки
    return s.strip()


# ---------- работа с API ----------
class RateLimitBlocked(Exception):
    """HDE ответил 429: аккаунт заблокирован примерно на 20 минут. Продолжать нельзя."""


class RateLimiter:
    """Скользящее окно в 60 с: не даёт отправить больше max_rpm запросов в минуту.

    Дополнительно читает заголовки лимита из ответа HDE (X-RateLimit-Remaining / X-RateLimit-Reset и
    их варианты), если сервер их присылает, и притормаживает, когда запас почти исчерпан.
    """
    REMAINING = ("X-RateLimit-Remaining", "X-Rate-Limit-Remaining", "RateLimit-Remaining")
    RESET = ("X-RateLimit-Reset", "X-Rate-Limit-Reset", "RateLimit-Reset")

    def __init__(self, rpm, safety):
        self.max = max(1, int(rpm * safety))
        self.hits = collections.deque()
        self.pause_until = 0.0

    def wait(self):
        while True:
            now = time.monotonic()
            if now < self.pause_until:
                time.sleep(self.pause_until - now)
                continue
            while self.hits and now - self.hits[0] >= 60:
                self.hits.popleft()
            if len(self.hits) < self.max:
                self.hits.append(now)
                return
            time.sleep(60 - (now - self.hits[0]) + 0.05)

    def note_headers(self, headers):
        def first(names):
            for n in names:
                v = headers.get(n)
                if v is not None:
                    return v
            return None
        try:
            remaining = first(self.REMAINING)
            if remaining is not None and int(remaining) <= 2:
                reset = first(self.RESET)
                delay = 60.0
                if reset is not None:
                    r = float(reset)
                    delay = r - time.time() if r > 1e9 else r      # unix-время или секунды
                self.pause_until = time.monotonic() + min(max(delay, 1.0), 60.0)
        except (TypeError, ValueError):
            pass


LIMITER = RateLimiter(RATE_LIMIT_RPM, RATE_SAFETY)


def api_get(path: str):
    for attempt in range(4):
        LIMITER.wait()
        try:
            r = requests.get(f"{BASE}{path}", auth=(HDE_EMAIL, HDE_API_KEY), timeout=30)
            LIMITER.note_headers(r.headers)
            if r.status_code == 200:
                return r.json()
            if r.status_code == 429:
                # Повторять нельзя: каждый лишний запрос — риск продлить блокировку на аккаунте.
                raise RateLimitBlocked(
                    f"HDE вернул 429: превышен лимит {RATE_LIMIT_RPM} запросов/мин, аккаунт заблокирован "
                    f"примерно на {BLOCK_MINUTES} минут. Подождите и запустите снова: уже скачанные заявки "
                    f"сохранены в кэше.")
            if r.status_code in (500, 502, 503, 504):
                time.sleep(2 * (attempt + 1))
                continue
            if r.status_code in (401, 403):
                return {"__error__": f"HTTP {r.status_code}: неверный email/ключ API или нет доступа к API"}
            return {"__error__": f"HTTP {r.status_code}"}
        except requests.RequestException:
            time.sleep(2 * (attempt + 1))
    return {"__error__": "сеть/таймаут"}


def extract_key(data) -> str:
    """Ищет id переписки (КСЭ.../b2b.../drr...) где угодно в JSON заявки."""
    blob = json.dumps(data, ensure_ascii=False)
    m = KEY_RE_ID.search(blob) or KEY_RE_FALLBACK.search(blob)
    return m.group(1) if m else ""


def get_posts(tid: str) -> list:
    """Все сообщения заявки: GET /tickets/{id}/posts/ (с учётом постраничности)."""
    posts, seen, page = [], set(), 1
    while page <= 50:
        data = api_get(f"/tickets/{tid}/posts/?page={page}")
        if "__error__" in data:
            raise RuntimeError(data["__error__"])
        chunk = data.get("data", data)
        if isinstance(chunk, dict):
            chunk = list(chunk.values())
        if not isinstance(chunk, list):
            break
        new = 0
        for p in chunk:
            if not isinstance(p, dict) or not any(k in p for k in ("message", "text", "body")):
                continue
            pid = p.get("id", id(p))
            if pid in seen:
                continue
            seen.add(pid)
            posts.append(p)
            new += 1
        pag = data.get("pagination") or {}
        total = pag.get("total_pages") or pag.get("pages") or pag.get("last_page") or 1
        if not new or page >= int(total):
            break
        page += 1
    return posts


def post_dt(p: dict) -> str:
    """Дата сообщения из полей поста (если в тексте нет метки [время дата])."""
    for k in ("date_created", "created_at", "date", "created"):
        v = p.get(k)
        if not v:
            continue
        if isinstance(v, (int, float)) or str(v).isdigit():
            try:
                return datetime.fromtimestamp(int(v), MSK).replace(tzinfo=None).isoformat()
            except (ValueError, OSError):
                continue
        for fmt in ("%H:%M:%S %d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d.%m.%Y %H:%M:%S"):
            try:
                return datetime.strptime(str(v).strip()[:19], fmt).isoformat()
            except ValueError:
                pass
    return ""


def parse_dt(msg_html: str) -> str:
    m = DATE_PARSE.search(msg_html[:40])
    if not m:
        return ""
    h, mi, d, mo, y = map(int, m.groups())
    try:
        return datetime(y, mo, d, h, mi).isoformat()
    except ValueError:
        return ""


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip().lower()


AUTHOR_KEYS = ("user", "author", "name", "from", "sender", "owner")


def author_info(post: dict) -> str:
    """Поля сообщения, похожие на автора (для распознавания ИИ и для --posts)."""
    parts = []
    for k, v in post.items():
        if any(a in k.lower() for a in AUTHOR_KEYS):
            parts.append(f"{k}=" + (json.dumps(v, ensure_ascii=False)
                                    if isinstance(v, (dict, list)) else str(v)))
    return "; ".join(parts)


# Запись вида «2026-10-20 06:22:39 Иванов упомянул(а) вас в запросе <ТЕКСТ ЗАПРОСА>»:
# заголовок отбрасываем, сам запрос клиента сохраняем.
MENTION_RE = re.compile(r".*?упомянул(?:\(а\)|а)?\s+вас\s+в\s+запросе[:\s]*", re.I | re.S)


def classify(msg_html: str, post: dict | None = None):
    text = clean_html(msg_html)
    m = MENTION_RE.match(text)
    if m:
        text = text[m.end():].strip()
    if not text:
        return "system", ""
    n = norm(text)
    if post is not None and str(post.get("user_id")) in AI_USER_IDS:
        return "ai", text
    marker = norm(AI_MARKER)
    if marker in n[:300] or (post and marker in norm(author_info(post))):
        return "ai", text
    if "департамент" in n:
        return "support", text
    return "client", text


def fetch_ticket(tid: str) -> dict:
    """Возвращает {'key': ..., 'msgs': [[iso_dt, kind, text], ...]} (тексты обезличены)."""
    data = api_get(f"/tickets/{tid}/")          # карточка заявки — для id переписки
    if "__error__" in data:
        raise RuntimeError(data["__error__"])
    posts = get_posts(tid)                       # сообщения заявки
    key = extract_key(data) or extract_key(posts)

    msgs = []
    for p in posts:
        body = p.get("message") or p.get("text") or p.get("body") or ""
        for part in DATE_RE.split(body):
            if not part.strip():
                continue
            kind, text = classify(part, p)
            if kind == "system" or not text:
                continue
            text = anonymize(text)
            if kind == "ai":
                text = text.replace(AI_MARKER, "").strip()
            msgs.append([parse_dt(part) or post_dt(p), kind, text])
    return {"key": key, "msgs": msgs}


# ---------- склейка по id переписки ----------
LABELS = {"client": "КЛИЕНТ", "ai": "ИИ", "support": "ПОДДЕРЖКА"}


def group_items(tids: list, cache: dict):
    """Сообщения всех заявок группы: сортировка по времени, без дублей. -> [(iso_dt, kind, text)]"""
    seen, items = set(), []
    order = 0
    for tid in tids:
        for dt, kind, text in cache[tid]["msgs"]:
            sig = (dt, kind, text[:200])
            if sig in seen:
                continue
            seen.add(sig)
            items.append((dt or "9999", order, kind, text))
            order += 1
    items.sort(key=lambda x: (x[0], x[1]))
    return [(("" if dt == "9999" else dt), k, t) for dt, _, k, t in items]


def date_from_name(name: str):
    """Дата из имени файла: queue_daily_20261025.xlsx -> date(2026, 10, 25). Нет даты -> None."""
    for m in re.finditer(r"(?<!\d)(\d{8})(?!\d)", name):
        try:
            return datetime.strptime(m.group(1), "%Y%m%d").date()
        except ValueError:
            continue
    return None


def filter_day(items, day):
    """Оставляет сообщения указанного дня. Сообщения без даты оставляются (их нельзя отнести к дню).
    Возвращает (items, число_без_даты)."""
    if day is None:
        return items, 0
    iso = day.isoformat()
    out = [(dt, k, t) for dt, k, t in items if not dt or dt[:10] == iso]
    return out, sum(1 for dt, _, _ in items if not dt)


def build_interactions(items):
    """Пары «сообщение клиента — ответ ИИ» в порядке переписки.
    Подряд идущие сообщения клиента (или ИИ) склеиваются. Сообщения поддержки в пары не входят."""
    pairs, client, ai = [], [], []
    for _, kind, text in items:
        if kind == "client":
            if ai:                       # клиент пишет снова после ответа ИИ — предыдущая пара закрыта
                pairs.append(("\n".join(client), "\n".join(ai)))
                client, ai = [], []
            client.append(text)
        elif kind == "ai":
            ai.append(text)
    if client or ai:
        pairs.append(("\n".join(client), "\n".join(ai)))
    return pairs


PAIR_FORMAT = "Клиент: {client}\nИИ: {ai}"       # формат одной пары в колонке «Запрос клиента»


def format_interactions(pairs):
    return "\n---\n".join(PAIR_FORMAT.format(client=c or "—", ai=a or "— (ответа ИИ нет)") for c, a in pairs)


def merge_group(items, day):
    """-> (запрос_клиента_по_дню, ответы_ии, диалог, client_text_дня, ai_text_дня, число_без_даты)"""
    reply_kinds = ("ai",) if ONLY_AI_REPLIES else ("ai", "support")
    day_items, undated = filter_day(items, day)
    pairs = build_interactions(day_items)
    client_day = "\n".join(t for _, k, t in day_items if k == "client")
    ai_day = "\n".join(t for _, k, t in day_items if k == "ai")
    replies = [t for _, k, t in items if k in reply_kinds]
    dialog = [f"[{LABELS.get(k, k)}] {t}" for _, k, t in items]
    cut = lambda x: x[:CELL_LIMIT]
    sep = "\n---\n"
    return (cut(format_interactions(pairs)), cut(sep.join(replies)), cut(sep.join(dialog)),
            client_day, ai_day, undated)


# ---------- Excel ----------
def find_input_file() -> Path:
    files = [f for f in INPUT_FOLDER.glob("*.xlsx") if not f.name.startswith("~$")]
    if not files:
        sys.exit(f"В папке {INPUT_FOLDER} нет .xlsx файлов")
    return max(files, key=lambda f: f.stat().st_mtime)


def save_wb(wb, path):
    try:
        wb.save(path)
    except PermissionError:
        sys.exit("Не могу сохранить файл — закройте его в Excel и запустите снова "
                 "(выгруженные заявки лежат в кэше, повторно качаться не будут).")


def pick_sheet(wb):
    """Активный лист, а если в нём нет колонки «Заявка» — первый лист, где она есть."""
    def has_id(w):
        return any(c.value and str(c.value).strip() == ID_COLUMN for c in w[1])
    if has_id(wb.active):
        return wb.active
    for w in wb.worksheets:
        if hde_rules.norm(w.title) != hde_rules.norm(hde_rules.ESC_SHEET) and has_id(w):
            return w
    return wb.active


def need_credentials():
    if not HDE_EMAIL or not HDE_API_KEY:
        sys.exit("Не заданы HDE_EMAIL и/или HDE_API_KEY. Задайте переменные окружения "
                 "в ЭТОМ ЖЕ окне терминала или впишите значения в настройки вверху файла.")


def run_report_mode(path, wb, source_name, day):
    """Готовый отчёт («Обращения» с колонками «Итог» и «Запрос клиента»): заявки скачивать не нужно.
    Только проверка правил и «Корректно». Сеть и ключ HDE не нужны."""
    print("Режим: готовый отчёт — заявки из HDE не скачиваются, запросов к API нет.")
    in_report = hde_rules.report_date(wb)
    if in_report and day and in_report.replace(" ", "") != day.strftime("%d.%m.%Y"):
        print(f"ВНИМАНИЕ: в имени файла дата {day.strftime('%d.%m.%Y')}, а на листе «Сводка» — {in_report}. "
              "Проверьте, тот ли файл загружен.")
    hde_rules.apply_report(wb)
    hde_rules.apply_escalations(wb)
    save_wb(wb, path)
    print("Готово.")


def main():
    if "--posts" in sys.argv:
        need_credentials()
        tid = sys.argv[sys.argv.index("--posts") + 1]
        posts = get_posts(tid)
        print(f"Сообщений: {len(posts)}  (id переписки: {extract_key(posts) or '—'})")
        for i, p in enumerate(posts, 1):
            body = p.get("message") or p.get("text") or p.get("body") or ""
            kind, text = classify(body, p)
            prev = anonymize(text)[:200].replace("\n", " ")
            print(f"{i:>2} | {post_dt(p) or '—'} | {kind:<7} | {anonymize(author_info(p))[:120]} | {prev}")
        print("Ключи сообщения:", sorted(posts[0].keys()) if posts else "—")
        return
    if "--debug" in sys.argv:
        need_credentials()
        tid = sys.argv[sys.argv.index("--debug") + 1]
        data = api_get(f"/tickets/{tid}/")
        if "__error__" in data:
            sys.exit("Ошибка запроса заявки: " + data["__error__"])
        posts = get_posts(tid)
        print("Найденный id переписки:", extract_key(data) or extract_key(posts) or "— не найден —")
        print("Сообщений в заявке:", len(posts))
        print("--- заявка ---")
        print(json.dumps(data, ensure_ascii=False, indent=2)[:3000])
        print("--- первые сообщения ---")
        print(json.dumps(posts[:3], ensure_ascii=False, indent=2)[:4000])
        return
    path = find_input_file()
    source_name = os.getenv("SOURCE_NAME") or path.name       # веб-версия копирует файл в input.xlsx
    day = date_from_name(source_name)
    print("Файл:", source_name)
    print(f"Дата переписок из имени файла: {day.strftime('%d.%m.%Y')}" if day else
          "В имени файла нет даты вида 20261025 — «Запрос клиента» собирается за всё время.")
    wb = load_workbook(path)
    if hde_rules.is_report(wb):
        return run_report_mode(path, wb, source_name, day)
    need_credentials()
    print("Режим: список номеров заявок — заявки будут скачаны из HDE.")
    ws = pick_sheet(wb)

    # --- 0) удалить колонки прошлых версий
    old = sorted((c.column for c in ws[1] if c.value and str(c.value).strip() in REMOVED_COLS), reverse=True)
    for idx in old:
        ws.delete_cols(idx)
    if old:
        print(f"Удалено колонок: {len(old)} ({', '.join(REMOVED_COLS)})")

    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    if ID_COLUMN not in headers:
        sys.exit(f"Не найдена колонка «{ID_COLUMN}». Есть: {list(headers)}")
    id_col = headers[ID_COLUMN]
    cols = {}
    for name in OUT_COLS:
        if name not in headers:
            headers[name] = ws.max_column + 1
            ws.cell(1, headers[name], name)
        cols[name] = headers[name]

    # --- 1) выгрузка заявок в кэш (при повторном запуске уже скачанные пропускаются)
    cache_path = path.with_suffix(".cache.json")
    cache = json.loads(cache_path.read_text(encoding="utf-8")) if cache_path.exists() else {}
    rows = {}  # row -> tid
    for row in range(2, ws.max_row + 1):
        v = ws.cell(row, id_col).value
        if v not in (None, ""):
            rows[row] = str(v).strip().lstrip("#")

    errors, fetched = {}, 0
    for row, tid in rows.items():
        if tid in cache:
            continue
        try:
            cache[tid] = fetch_ticket(tid)
        except RateLimitBlocked as e:
            cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
            sys.exit(str(e))
        except Exception as e:
            errors[tid] = str(e)
        fetched += 1
        print(f"[{fetched}] заявка {tid}" + (f"  ОШИБКА: {errors[tid]}" if tid in errors else ""))
        if fetched % SAVE_EVERY == 0:
            cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
        time.sleep(PAUSE_SEC)
    cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")

    # --- 2) группировка по id переписки (нет id — заявка остаётся сама по себе)
    groups, first_row = {}, {}
    for row, tid in rows.items():
        if tid not in cache:
            continue
        gkey = cache[tid]["key"] or f"заявка-{tid}"
        groups.setdefault(gkey, [])
        if tid not in groups[gkey]:
            groups[gkey].append(tid)
        first_row.setdefault(gkey, row)

    # --- 3) запись в Excel
    ctx, undated_total, empty_day = {}, 0, 0      # ctx[gkey] = (текст клиента за день, текст ИИ за день)
    for row, tid in rows.items():
        if tid in errors:
            ws.cell(row, cols["Запрос клиента"], f"Ошибка: {errors[tid]}")
            continue
        if tid not in cache:
            continue
        gkey = cache[tid]["key"] or f"заявка-{tid}"
        ws.cell(row, cols["ID переписки"], cache[tid]["key"])
        if first_row[gkey] == row:
            q, a, d, client_day, ai_day, undated = merge_group(group_items(groups[gkey], cache), day)
            ctx[gkey] = (client_day, ai_day)
            undated_total += undated
            empty_day += 0 if q else 1
            ws.cell(row, cols["Запрос клиента"], q)
            ws.cell(row, cols["Ответы ИИ"], a)
            ws.cell(row, cols["Диалог (размечено)"], d)
        else:
            for c in OUT_COLS[1:]:
                ws.cell(row, cols[c]).value = None      # ws.cell(.., None) НЕ очищает ячейку
    if day and undated_total:
        print(f"Внимание: {undated_total} сообщений без даты оставлены в «Запрос клиента» (день определить нельзя).")
    if day and empty_day:
        print(f"За {day.strftime('%d.%m.%Y')} нет сообщений в {empty_day} группах — «Запрос клиента» пустой.")

    hde_rules.apply_escalations(wb)

    save_wb(wb, path)
    merged = sum(1 for g in groups.values() if len(g) > 1)
    print(f"Готово. Групп с несколькими заявками: {merged}. Без id переписки: "
          f"{sum(1 for t in cache.values() if not t['key'])}. Ошибок: {len(errors)}")


if __name__ == "__main__":
    main()
