import sys
import unittest
from datetime import datetime, time as dtime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))
import hde_rules as r  # noqa: E402
from openpyxl import Workbook  # noqa: E402

RET = "Возврат оформлен + ответ клиенту"
PROOF = "Запрос подтверждения отмены → ответ клиенту"
D_OK = "Дата доставки изменена + ответ клиенту"
D_UPD = "Дата уже стоит в УПД → ответ клиенту по шаблону"
D_NA = "Дата недоступна → клиенту отправлен список вариантов"
OPER = "Не удалось автоматически → оператору"


def v(client="", subject="", result="", note=""):
    return r.report_verdict(client, subject, note, result)


class ReportRules(unittest.TestCase):
    # ---- отмена
    def test_cancel_words_in_client_text(self):
        for c in ("Просьба его отменить", "Отмена заказа", "Просим вернуть посылку на склад",
                  "клиент отказался от заказа, просьба вернуть посылку", "Клиент Отказался",
                  "Просьба поставить заказ на возврат", "заказ больше не актуален, не актуальна доставка",
                  "Клиент передумал"):
            self.assertEqual(v(client=c, result=RET), (True, "отмена"), c)

    def test_cancel_words_only_in_subject_column_d(self):
        self.assertEqual(v(client="Добрый день, просьба оформить", subject="Отказ клиента", result=RET),
                         (True, "отмена"))
        self.assertEqual(v(client="Добрый день", subject="Sunlight: Отмена заказа: 1", result=RET),
                         (True, "отмена"))

    def test_cancel_old_phrases_still_work(self):
        self.assertTrue(v(client="Отмена", result="Отменена часть списка → остальное оператору")[0])
        self.assertTrue(v(client="Возврат", result="Отменена часть списка -> остальное оператору")[0])
        self.assertTrue(v(client="Отмена", result=OPER,
                          note="Информация о необходимости возврата внесена в накладную.")[0])
        self.assertTrue(v(client="Отмена", result=OPER,
                          note="По данному отправлению возврат уже оформлен ранее, повторная заявка не требуется")[0])

    def test_cancel_negative(self):
        self.assertFalse(v(client="Где моя посылка?", subject="Статус", result=RET)[0])   # клиент не просит отмену
        self.assertFalse(v(client="Отмена", result=OPER)[0])                              # ИИ не оформил
        self.assertFalse(v(client="Отмена", result="Отмена невозможна → клиенту дан ответ + оператору")[0])
        self.assertFalse(v(client="", subject="", result=RET)[0])

    # ---- подтверждение отмены
    def test_proof_needs_only_f(self):
        self.assertEqual(v(result=PROOF), (True, "подтверждение отмены"))
        self.assertEqual(v(client="любой текст", result="запрос подтверждения отмены -> ответ клиенту")[0], True)
        self.assertFalse(v(result="Подтверждение просят, а отмены не было → оператору")[0])

    # ---- перенос даты
    def test_date_change(self):
        for c in ("Простьба доставить 30.10.2026", "Просьба перенести доставку на 02.11",
                  "Просьба изменить дату доставки 19698663177 Трекномер 497-A05",
                  "Изменить дату доставки на пятницу", "Перенесите доставку, пожалуйста",
                  "Согласован перенос на 28.10 интервал 8.00-23.00", "доставить 31.10"):
            for f in (D_OK, D_UPD, D_NA):
                self.assertEqual(v(client=c, result=f), (True, "перенос даты"), (c, f))

    def test_date_change_from_subject(self):
        self.assertTrue(v(client="Добрый день", subject="SDIT-1 - Заявка - Изменить дату доставки - 1",
                          result=D_UPD)[0])

    def test_date_change_negative(self):
        self.assertFalse(v(client="Где мой груз?", result=D_OK)[0])
        self.assertFalse(v(client="Просьба доставить 30.10", result=OPER)[0])
        self.assertFalse(v(client="Просьба доставить 30.10", result="")[0])

    def test_none_values(self):
        self.assertEqual(r.report_verdict(None, None, None, None), (False, ""))

    def test_priority_proof_before_others(self):
        self.assertEqual(v(client="Отмена", result=PROOF)[1], "подтверждение отмены")


def make_report(order=("Заявка", "Время", "Тип запроса", "Тема", "Накладная", "Итог", "Запрос клиента",
                       "Причина / примечание", "Комментарий КСЭ")):
    wb = Workbook()
    ws = wb.active
    ws.title = "Обращения"
    ws.append(list(order))
    data = {
        "#1": dict(Тема="Sunlight: Отмена заказа", Итог=RET, **{"Запрос клиента": "просьба отменить"}),
        "#2": dict(Тема="x", Итог=D_OK, **{"Запрос клиента": "Просьба доставить 30.10"}),
        "#3": dict(Тема="x", Итог=PROOF, **{"Запрос клиента": "подтвердите"}),
        "#4": dict(Тема="x", Итог=OPER, **{"Запрос клиента": "что-то"}),
        "#5": dict(Тема="x", Итог=RET, **{"Запрос клиента": "Отмена", "Комментарий КСЭ": "моё мнение"}),
        "#6": dict(Тема="x", Итог=RET, **{"Запрос клиента": "Отмена", "Комментарий КСЭ": "Корректно"}),
    }
    for tid, d in data.items():
        d["Заявка"] = tid
        ws.append([d.get(h) for h in order])
    return wb, ws


class ReportSheet(unittest.TestCase):
    def test_apply_report(self):
        wb, ws = make_report()
        logs = []
        st = r.apply_report(wb, logs.append)
        col = 9
        self.assertEqual([ws.cell(i, col).value for i in range(2, 8)],
                         [r.OK, r.OK, r.OK, None, "моё мнение", r.OK])
        self.assertEqual((st["marked"], st["occupied"]), (3, 1))
        self.assertEqual(st["left"], {OPER: 1})

    def test_columns_found_by_header_not_letter(self):
        wb, ws = make_report(order=("Комментарий КСЭ", "Запрос клиента", "Итог", "Тема", "Заявка"))
        r.apply_report(wb, lambda m: None)
        self.assertEqual([ws.cell(i, 1).value for i in range(2, 8)],
                         [r.OK, r.OK, r.OK, None, "моё мнение", r.OK])

    def test_missing_opinion_column_is_added(self):
        wb, ws = make_report(order=("Заявка", "Тема", "Итог", "Запрос клиента"))
        r.apply_report(wb, lambda m: None)
        self.assertEqual(ws.cell(1, 5).value, "Комментарий КСЭ")
        self.assertEqual(ws.cell(2, 5).value, r.OK)

    def test_is_report(self):
        wb, _ = make_report()
        self.assertTrue(r.is_report(wb))
        legacy = Workbook()
        legacy.active.append(["Заявка"])
        self.assertFalse(r.is_report(legacy))
        no_cols = Workbook()
        no_cols.active.title = "Обращения"
        no_cols.active.append(["Заявка", "Тема"])
        self.assertFalse(r.is_report(no_cols))

    def test_rerun_is_idempotent(self):
        wb, ws = make_report()
        r.apply_report(wb, lambda m: None)
        before = [[c.value for c in row] for row in ws.iter_rows()]
        r.apply_report(wb, lambda m: None)
        self.assertEqual(before, [[c.value for c in row] for row in ws.iter_rows()])

    def test_report_date(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Сводка"
        ws.append(["Дата", "27.10.2026"])
        self.assertEqual(r.report_date(wb), "27.10.2026")
        self.assertIsNone(r.report_date(Workbook()))


class Escalations(unittest.TestCase):
    def test_domain(self):
        self.assertEqual(r.email_domain("test@skx.cse.ru"), "skx.cse.ru")
        self.assertEqual(r.email_domain("Иван <Test2@SKX.cse.ru>"), "skx.cse.ru")
        self.assertIsNone(r.email_domain("нет почты"))
        self.assertIsNone(r.email_domain(None))

    def test_to_domains_ignores_copy(self):
        self.assertEqual(r.to_domains("Кому: office@pkc.cse.ru"), {"pkc.cse.ru"})
        self.assertEqual(r.to_domains("Кому: logistics@tya.cse.ru\nКопия: a@tya.cse.ru, b@vgd.cse.ru"),
                         {"tya.cse.ru"})
        self.assertEqual(r.to_domains("Кому: a@x.ru, b@y.ru\nКопия: c@z.ru"), {"x.ru", "y.ru"})
        self.assertEqual(r.to_domains("test@skx.cse.ru"), {"skx.cse.ru"})     # без слова «Кому»
        self.assertEqual(r.to_domains(None), set())

    def test_copy_domain_does_not_count(self):
        ok, why = r.escalation_verdict("10:00", "Москва", "Кому: a@x.ru\nКопия: b@vgd.cse.ru",
                                       "Управляющий офисом", "boss@vgd.cse.ru", "Эскалация выполнена")
        self.assertFalse(ok)
        self.assertIn("домены", why)

    def test_rule1_same_domain(self):
        ok, why = r.escalation_verdict("10:00", "Москва", "test@skx.cse.ru", "Управляющий офисом",
                                       "test2@skx.cse.ru", "В работе")
        self.assertEqual((ok, why), (True, "правило 1"))

    def test_rule1_different_domain(self):
        ok, _ = r.escalation_verdict("10:00", "Москва", "a@skx.cse.ru", "Управляющий офисом", "b@other.ru", "")
        self.assertFalse(ok)

    def test_not_manager(self):
        for role in ("Региональный управляющий", "ОКОВРС", "ОКОВРС вместо рег. управляющего"):
            ok, _ = r.escalation_verdict("10:00", "Москва", "a@x.ru", role, "b@x.ru", "")
            self.assertFalse(ok, role)

    def test_rule2_chelyabinsk_example(self):
        # B=17:18 (UTC+3), Челябинск (+5) → 19:18 → вне 9–18 → корректно
        ok, why = r.escalation_verdict("17:18", "КС Челябинск", "Кому: a@skx.cse.ru", "Управляющий офисом",
                                       "b@skx.cse.ru", "Ждёт рабочего времени офиса")
        self.assertEqual((ok, why), (True, "правило 2"))

    def test_rule2_inside_hours(self):
        ok, why = r.escalation_verdict("12:00", "Челябинск", "a@skx.cse.ru", "Управляющий офисом",
                                       "b@skx.cse.ru", "Ждёт рабочего времени офиса")
        self.assertFalse(ok)
        self.assertIn("рабочее", why)

    def test_rule2_boundaries(self):
        f = lambda t, reg: r.escalation_verdict(t, reg, "a@x.ru", "Управляющий офисом", "b@x.ru",
                                                "ждет рабочего времени офиса")[0]
        self.assertFalse(f("09:00", "Москва"))
        self.assertTrue(f("08:59", "Москва"))
        self.assertTrue(f("18:00", "Москва"))
        self.assertFalse(f("17:59", "Москва"))
        self.assertTrue(f("23:30", "Владивосток"))
        self.assertFalse(f("03:00", "Владивосток"))
        self.assertTrue(f("08:00", "Калининград"))

    def test_rule2_unknown_region_and_bad_time(self):
        a = ("a@x.ru", "Управляющий офисом", "b@x.ru", "Ждёт рабочего времени офиса")
        ok, why = r.escalation_verdict("17:00", "Атлантида", *a)
        self.assertFalse(ok)
        self.assertIn("неизвестный регион", why)
        self.assertFalse(r.escalation_verdict("не время", "Москва", *a)[0])

    def test_time_parsing(self):
        self.assertEqual(r.parse_time("29.10.2026 17:18"), dtime(17, 18))
        self.assertEqual(r.parse_time(datetime(2026, 10, 29, 9, 5)), dtime(9, 5))
        self.assertEqual(r.parse_time(dtime(7, 0)), dtime(7, 0))
        self.assertEqual(r.parse_time(0.75), dtime(18, 0))
        self.assertIsNone(r.parse_time(""))
        self.assertIsNone(r.parse_time("25:99"))

    def test_regions(self):
        cases = {"Челябинск": 5, "КС Челябинск": 5, "Челябинская область": 5, "Томск": 7, "Омск": 6,
                 "Москва": 3, "Нижневартовск": 5, "Нижний Новгород": 3, "Киров": 3, "Кировград": None,
                 "Уфа": 5, "Сахалинская область": 11, "Республика Саха (Якутия)": 9, "Чита": 9, "": None,
                 # добавлено по отчёту 27.10.2026
                 "КС Пятигорск": 3, "КС Грозный": 3, "КС Нальчик": 3, "КС Магнитогорск": 5,
                 "КС Норильск": 7, "КС Ангарск": 8, "КС Братск": 8, "КС Артем": 10, "КС Артём": 10,
                 "КС Комсомольск-на-Амуре": 10, "КС Находка": 10, "КС Петропавловск-Камчатский": 12,
                 "КС Южно-Сахалинск": 11, "КС Ростов-на-Дону": 3, "КС Владивосток": 10}
        for reg, off in cases.items():
            self.assertEqual(r.region_offset(reg), off, reg)

    def test_apply_sheet_by_headers(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Эскалации"
        ws.append(["Заявка", "Время", "Тема", "Офис", "Изначальные адресаты", "Этап до эскалации",
                   "Кого подключили", "Адрес", "Итог", "Примечание", "Комментарий КСЭ"])
        W = "Ждёт рабочего времени офиса"
        rows = [
            ("#1", "10:00", "t", "КС Москва", "Кому: a@skx.cse.ru", "", "Управляющий офисом", "b@skx.cse.ru", "Эскалация выполнена", "", None),   # правило 1
            ("#2", "10:00", "t", "КС Москва", "Кому: a@skx.cse.ru", "", "Управляющий офисом", "b@zzz.ru", "Эскалация выполнена", "", None),       # домен
            ("#3", "17:18", "t", "КС Челябинск", "Кому: a@skx.cse.ru", "", "Управляющий офисом", "b@skx.cse.ru", W, "", None),                    # правило 2
            ("#4", "12:00", "t", "КС Челябинск", "Кому: a@skx.cse.ru", "", "Управляющий офисом", "b@skx.cse.ru", W, "", None),                    # рабочее время
            ("#5", "10:00", "t", "КС Москва", "Кому: a@skx.cse.ru", "", "ОКОВРС", "okovrs@cse.ru", "Эскалация выполнена", "", None),              # не управляющий
            ("#6", "10:00", "t", "КС Москва", "Кому: a@skx.cse.ru", "", "Управляющий офисом", "b@skx.cse.ru", "Эскалация выполнена", "", "Ручная пометка"),
            ("#7", "15:02", "t", "КС Ангарск", "Кому: a@ang.cse.ru", "", "Управляющий офисом", "b@ang.cse.ru", W, "", None),                      # 20:02 местного
        ]
        for row in rows:
            ws.append(row)
        logs = []
        st = r.apply_escalations(wb, logs.append)
        self.assertEqual([ws.cell(i, 11).value for i in range(2, 9)],
                         [r.OK, None, r.OK, None, None, "Ручная пометка", r.OK])
        self.assertEqual((st["marked"], st["rule1"], st["rule2"], st["occupied"]), (3, 1, 2, 1))
        self.assertEqual(st["other_roles"], {"ОКОВРС": 1})

    def test_missing_sheet_and_missing_columns(self):
        self.assertEqual(r.apply_escalations(Workbook(), lambda m: None), {"marked": 0})
        wb = Workbook()
        wb.active.title = "Эскалации"
        wb.active.append(["Заявка", "Время"])
        self.assertEqual(r.apply_escalations(wb, lambda m: None), {"marked": 0})


if __name__ == "__main__":
    unittest.main()
