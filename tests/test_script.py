"""Offline tests for hde_ai_analysis.py: rate limiter, 429 handling, date from file name, pairs."""
import os
import sys
import unittest
from datetime import date
from pathlib import Path

os.environ.setdefault("HDE_EMAIL", "x")
os.environ.setdefault("HDE_API_KEY", "y")
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))
import hde_ai_analysis as h  # noqa: E402


class FakeClock:
    def __init__(self):
        self.t = 1000.0
        self._real = (h.time.monotonic, h.time.sleep, h.time.time)

    def __enter__(self):
        h.time.monotonic = lambda: self.t
        h.time.sleep = lambda s: setattr(self, "t", self.t + s)
        h.time.time = lambda: 1.8e9 + self.t
        return self

    def __exit__(self, *a):
        h.time.monotonic, h.time.sleep, h.time.time = self._real


class Resp:
    def __init__(self, code, headers=None):
        self.status_code, self.headers = code, headers or {}

    def json(self):
        return {"ok": 1}


class Limiter(unittest.TestCase):
    def test_window(self):
        with FakeClock() as c:
            lim = h.RateLimiter(300, 0.9)
            self.assertEqual(lim.max, 270)
            for _ in range(270):
                lim.wait()
            self.assertEqual(c.t, 1000.0)          # first 270 go through immediately
            lim.wait()
            self.assertGreaterEqual(c.t - 1000.0, 60.0)

    def test_never_over_limit(self):
        with FakeClock() as c:
            lim, stamps = h.RateLimiter(300, 0.9), []
            for _ in range(800):
                lim.wait()
                stamps.append(c.t)
            self.assertLessEqual(max(sum(1 for s in stamps if t <= s < t + 60) for t in stamps), 270)

    def test_429_stops_without_retry(self):
        calls = []
        orig = h.requests.get
        h.requests.get = lambda *a, **k: (calls.append(1), Resp(429))[1]
        h.LIMITER = h.RateLimiter(300, 0.9)
        try:
            with self.assertRaises(h.RateLimitBlocked):
                h.api_get("/x")
        finally:
            h.requests.get = orig
        self.assertEqual(len(calls), 1)

    def test_remaining_header_slows_down(self):
        orig = h.requests.get
        h.requests.get = lambda *a, **k: Resp(200, {"X-RateLimit-Remaining": "1", "X-RateLimit-Reset": "30"})
        try:
            with FakeClock() as c:
                h.LIMITER = h.RateLimiter(300, 0.9)
                h.api_get("/a")
                h.api_get("/b")
                self.assertAlmostEqual(c.t - 1000.0, 30.0, delta=1)
        finally:
            h.requests.get = orig


class Dates(unittest.TestCase):
    def test_date_from_name(self):
        self.assertEqual(h.date_from_name("queue_daily_20261025.xlsx"), date(2026, 10, 25))
        self.assertEqual(h.date_from_name("20261025.xlsx"), date(2026, 10, 25))
        self.assertIsNone(h.date_from_name("queue_daily.xlsx"))
        self.assertIsNone(h.date_from_name("report_20261345.xlsx"))      # нет такого месяца
        self.assertIsNone(h.date_from_name("id_123456789012.xlsx"))      # 12 цифр — не дата

    def test_pairs_and_day_filter(self):
        d1, d2 = "2026-10-24T10:00:00", "2026-10-25T10:00:00"
        items = [(d1, "client", "вчера"), (d1, "ai", "ответ вчера"),
                 (d2, "client", "A"), (d2, "client", "B"), (d2, "ai", "ответ AB"),
                 (d2, "client", "C"), ("", "client", "без даты")]
        day_items, undated = h.filter_day(items, date(2026, 10, 25))
        self.assertEqual(undated, 1)
        self.assertEqual(h.build_interactions(day_items),
                         [("A\nB", "ответ AB"), ("C\nбез даты", "")])
        out = h.format_interactions(h.build_interactions(day_items))
        self.assertIn("Клиент: A\nB\nИИ: ответ AB", out)
        self.assertNotIn("вчера", out)


if __name__ == "__main__":
    unittest.main()


class ReportMode(unittest.TestCase):
    """Готовый отчёт: без ключа HDE и без единого сетевого запроса."""

    def test_report_mode_offline(self):
        import tempfile
        from openpyxl import Workbook, load_workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Обращения"
        ws.append(["Заявка", "Тема", "Итог", "Запрос клиента", "Комментарий КСЭ"])
        ws.append(["#1", "Отмена заказа", "Возврат оформлен + ответ клиенту", "просьба отменить", None])
        ws.append(["#2", "x", "Не удалось автоматически → оператору", "что-то", None])
        s = wb.create_sheet("Сводка")
        s.append(["Дата", "27.10.2026"])
        calls = []
        orig = h.requests.get
        h.requests.get = lambda *a, **k: calls.append(a) or (_ for _ in ()).throw(AssertionError("network!"))
        saved = (h.HDE_EMAIL, h.HDE_API_KEY)
        h.HDE_EMAIL = h.HDE_API_KEY = ""                    # ключа нет — режим отчёта не должен его требовать
        try:
            with tempfile.TemporaryDirectory() as d:
                path = Path(d) / "input.xlsx"
                h.run_report_mode(path, wb, "queue_daily_20261027.xlsx", date(2026, 10, 27))
                out = load_workbook(path)["Обращения"]
                self.assertEqual([out.cell(2, 5).value, out.cell(3, 5).value], ["Корректно", None])
        finally:
            h.requests.get = orig
            h.HDE_EMAIL, h.HDE_API_KEY = saved
        self.assertEqual(calls, [])
